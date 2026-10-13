"""Document extractors (P1-T02): PDF and DOCX text + table extraction.

Tables are rendered as Markdown so they flow into chunks as plain text.
Optional dependencies live in ingestion/requirements.txt; a clear error
is raised if a needed one is missing.
"""

from __future__ import annotations

from pathlib import Path

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None


def table_to_markdown(table: list[list[str]]) -> str:
    """Render a table (list of rows of cell strings) as a Markdown table."""
    rows = [[(cell or "").strip().replace("\n", " ") for cell in row] for row in table]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join("---" for _ in rows[0]) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def extract_pdf(path: Path) -> str:
    """Extract text and tables from a PDF, page by page."""
    if pdfplumber is None:
        raise RuntimeError(
            "pdfplumber is not installed: pip install -r ingestion/requirements.txt"
        )
    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            parts.append(f"--- page {i} ---")
            parts.append(page.extract_text() or "")
            for table in page.extract_tables():
                md = table_to_markdown(table)
                if md:
                    parts.append(md)
    return "\n\n".join(part for part in parts if part.strip())


def extract_docx(path: Path) -> str:
    """Extract paragraphs and tables from a DOCX file."""
    if docx is None:
        raise RuntimeError(
            "python-docx is not installed: pip install -r ingestion/requirements.txt"
        )
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        md = table_to_markdown([[c.text for c in row.cells] for row in table.rows])
        if md:
            parts.append(md)
    return "\n\n".join(parts)
